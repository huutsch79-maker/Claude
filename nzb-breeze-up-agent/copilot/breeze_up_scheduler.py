import openpyxl,random,re,time
F='INPUT.xlsx'
GAP=4;WS=150;LANES=[7,8,10];SHIFT=2;RUNS=150;TRIES=3;SEED=26;BUDGET=60;VERS=3;DAYS=['Mon','Tue'];JW=15;JPEN=300;PRIO=[]
PAIRS={('prima park','mon'):2}
ALIAS={'mark brooks / alex olivera':'mark brooks'}
NOJ={'','no jockey','tbc','tba','n/a','-','none'}
def k(v):return re.sub(r'\s+',' ',str(v or '')).strip().lower()
wb=openpyxl.load_workbook(F);T0=time.time()
def rows(n):
 ws=next((w for w in wb if k(w.title)==k(n)),None)
 return [list(r) for r in ws.iter_rows(values_only=True)] if ws else []
def col(h,*ns):
 for n in ns:
  for i,x in enumerate(h):
   if k(x)==k(n):return i
 return -1
E=[];seen={};JID={};MS=[]
def day(d):
 R=rows(d);h=R[0];P=col(h,'Preparer Name','Preparer');B=col(h,'BUO');L=col(h,'Lot');J=col(h,'Jockey','Jockey #1','Rider');V=col(h,'Vendor/Draft Name','Vendor');G=col(h,'Breeding');D=col(h,'Day')
 H={};nm={}
 for r in R[1:]:
  if not r[P]:continue
  p=ALIAS.get(k(r[P]),k(r[P]));nm.setdefault(p,str(r[P]).strip());j=k(r[J]);lot=r[L]
  if lot in seen:E.append(f'ERROR lot {lot} listed twice ({seen[lot]} and {d})')
  seen[lot]=d
  if j in NOJ:E.append(f'INFO {d} lot {lot} has no jockey')
  if D>=0 and r[D] and k(r[D])!=k(d):E.append(f'WARNING lot {lot} is on the {d} sheet but Day says {r[D]}')
  H.setdefault((p,r[B]),[]).append((r[P],r[V] if V>=0 else '',r[B],lot,r[G] if G>=0 else '',str(r[J] or 'No Jockey').strip(),'' if j in NOJ else JID.setdefault(j,len(JID))))
 by={}
 for (p,b),hs in sorted(H.items(),key=lambda x:(x[0][0],float(x[0][1]))):by.setdefault(p,[]).append(hs)
 O0=[ALIAS.get(k(r[0]),k(r[0])) for r in rows(d+' Order')[1:] if r and r[0]];O=[p for p in O0 if p in by]+[p for p in by if p not in O0];U={}
 for pi,p in enumerate(O):
  if p not in O0:E.append(f'WARNING {d}: {nm[p]} not on {d} Order - scheduled last')
  n=PAIRS.get((p,k(d)),1);hs=by[p]
  for x in hs:
   js=[y[6] for y in x if y[6]!='']
   if len(js)>len(set(js)):E.append(f'ERROR {d}: {nm[p]} BUO {x[0][2]} has the same jockey on both horses')
  B2=[]
  for i in range(0,len(hs),n):
   g=hs[i:i+n];js=[y[6] for x in g for y in x if y[6]!='']
   if len(g)>1 and len(js)>len(set(js)):E.append(f'WARNING {d}: {nm[p]} BUO {g[0][0][2]}-{g[-1][0][2]} share a jockey, sent separately');B2+=[[x] for x in g]
   else:B2.append(g)
  U[p]=[(pi,r,g,[(o,y[6]) for o,x in enumerate(g) for y in x if y[6]!=''],len(g),SHIFT//n) for r,g in enumerate(B2)]
 for p in O:
  last={};t=-1
  for x in U[p]:
   t=max([t+1]+[last[j]+GAP+1-o for o,j in x[3] if j in last]);last.update({j:t+o for o,j in x[3]});t+=x[4]-1
  MS.append(t+1)
 return [U[p] for p in O]
def score(seq,NP,lanes):
 last=[-99]*len(JID);v=0;first=[-1]*NP;lp=[0]*NP;cnt=[0]*NP;mv=i=0
 for pi,r,u,J,n,S in seq:
  if first[pi]<0:first[pi]=i
  d=abs(r-cnt[pi]);mv+=d+(2e6 if d>S else 0);cnt[pi]+=1
  for o,j in J:
   e=i+o-last[j]
   if e<=GAP:v+=1e6+(GAP+1-e)*1e4
   elif last[j]>=0:w=GAP+2 if j in PJ else JW;v+=max(0,e-w)*JPEN*(1+2*(j in PJ))
   last[j]=i+o
  i+=n;lp[pi]=i-1
 return v+WS*sum(max(0,lp[p]-first[p]+1-MS[p])**2 for p in range(NP))+mv*20+100*sum(abs(a-b) for a,b in enumerate(sorted(range(NP),key=first.__getitem__)))
def greedy(U,lanes,rnd,M=1):
 NP=len(U);q=sorted(range(NP),key=lambda p:p+rnd.random()*2.9) if rnd else list(range(NP))
 act=[];last=[-99]*len(JID);seq=[];t=0;sv={};ad={}
 def ok(J):return all(t+o-last[j]>GAP for o,j in J) and len({j for o,j in J})==len(J)
 def add(p):act.append([p,list(U[p]),0]);sv[p]=t;ad[p]=len(ad)
 while act or q:
  while (not act or not M and len(act)<lanes) and q:add(q.pop(0))
  act.sort(key=lambda a:(ad if M else sv)[a[0]]+(rnd.random()*1.5 if rnd else 0));pick=None
  for a in act:
   S=a[1][0][5];lim=0 if a[2]-a[1][0][1]>=S else min(S,len(a[1])-1)
   c=[i for i in range(lim+1) if ok(a[1][i][3])]
   if c:pick=(a,c[0]);break
  while not pick and q and len(act)<lanes:
   add(q.pop(0))
   if ok(act[-1][1][0][3]):pick=(act[-1],0)
  if not pick:pick=min(((a,i) for a in act for i in range(min(a[1][0][5],len(a[1])-1)+1)),key=lambda x:sum(max(0,GAP+1-t-o+last[j]) for o,j in x[0][1][x[1]][3]))
  a,i=pick;x=a[1].pop(i);seq.append(x)
  for o,j in x[3]:last[j]=t+o
  t+=x[4];sv[a[0]]=t;a[2]+=1
  if not a[1]:act.remove(a)
 return seq
def polish(s,NP,lanes,rnd,end):
 c=b=score(s,NP,lanes);bs=s[:];N=max(12000,120*len(s));n=0
 while n<N and (n%200 or time.time()<end):
  T=3000*(1/1500)**(n/N);i=rnd.randrange(len(s));j=min(len(s)-1,max(0,i+rnd.randint(-2*lanes,2*lanes)));t=s[:];n+=1
  if rnd.random()<.5:t[i],t[j]=t[j],t[i]
  else:t.insert(j,t.pop(i))
  x=score(t,NP,lanes)
  if x<=c or rnd.random()<2.718**((c-x)/T):
   s,c=t,x
   if x<b:b,bs=x,t[:]
 return bs,b
out=openpyxl.load_workbook(F);summ=[];DD=[d for d in DAYS if rows(d)]
for di,dn in enumerate(DD):
 U=day(dn);NP=len(U);PJ={JID[k(n)] for n in PRIO if k(n) in JID};best=None;DL=time.time()+(T0+BUDGET-time.time())/(len(DD)-di)
 for vs in range(VERS):
  rnd=random.Random(SEED+vs);dl=time.time()+(DL-time.time())/(VERS-vs);bv=None
  for li,lanes in enumerate(LANES):
   le=dl if li==len(LANES)-1 else time.time()+(dl-time.time())/2
   for c in sorted((greedy(U,l2,rnd if x>1 else None,x%2) for x in range(RUNS) for l2 in ([lanes] if x%2 else [5,6,7])),key=lambda s:score(s,NP,lanes))[:TRIES]:
    s,b=polish(c,NP,lanes,rnd,time.time()+(le-time.time())/TRIES)
    if not bv or b<bv[1]:bv=(s,b,lanes)
   if bv[1]<1e6:break
  if not best or bv[1]<best[1]:best=bv
 seq,b,lanes=best;ws=out.create_sheet(f'{dn} Schedule')
 ws.append(['Heat','Preparer','Vendor','BUO','Lot','Breeding','Jockey','Heats since last ride','Check']);last={};cl=0;i=0;fh={};jr={};jn={}
 for pi,r,u,J,n,S in seq:
  for h in u:
   i+=1
   for x in h:
    g=i-last[x[6]]-1 if x[6] in last else '';bad=g!=''and g<GAP;cl+=bad;fh.setdefault(x[0],[]).append(i)
    ws.append([i,x[0],x[1],x[2],x[3],x[4],x[5],g,'CLASH' if bad else 'OK'])
    if x[6]!='':last[x[6]]=i;jr.setdefault(x[6],[]).append(i);jn.setdefault(x[6],x[5])
 jw=out.create_sheet(f'{dn} Jockeys');jw.append(['Jockey','Rides','First heat','Last heat','Longest wait (heats)',f'Waits over {JW} heats']);lw=0
 for j,hh in sorted(jr.items(),key=lambda z:z[1][0]-z[1][-1]):gp=[b-a-1 for a,b in zip(hh,hh[1:])];lw+=sum(g>JW for g in gp);jw.append([jn[j],len(hh),hh[0],hh[-1],max(gp,default=0),sum(g>JW for g in gp)])
 summ+=[f'{dn}: {i} heats, {cl} jockey clashes, max {lanes} preparers open, {lw} long jockey waits']+[f'  {dn} #{n+1} {q}: heats {min(v)}-{max(v)}' for n,(q,v) in enumerate(fh.items())]
v=out.create_sheet('Validation');[v.append([m]) for m in (summ+E or ['No E'])]
out.save('NZB_Breeze_Up_Schedule.xlsx');print('\n'.join([x for x in summ if x[0]!=' ']+[x for x in E if not x.startswith('INFO')]))
