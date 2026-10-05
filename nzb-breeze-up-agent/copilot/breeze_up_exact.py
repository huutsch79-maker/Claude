import openpyxl,re,time
from ortools.sat.python import cp_model as cp
F='INPUT.xlsx'
GAP=4;VMIN=4;VMAX=8;SHIFT=2;TL=90;DAYS=['Mon','Tue'];JW=15
PAIRS={('prima park','mon'):2}
ALIAS={'mark brooks / alex olivera':'mark brooks'}
NOJ={'','no jockey','tbc','tba','n/a','-','none'}
def k(v):return re.sub(r'\s+',' ',str(v or '')).strip().lower()
wb=openpyxl.load_workbook(F)
def rows(n):
 ws=next((w for w in wb if k(w.title)==k(n)),None)
 return [list(r) for r in ws.iter_rows(values_only=True)] if ws else []
def col(h,*ns):
 for n in ns:
  for i,x in enumerate(h):
   if k(x)==k(n):return i
 return -1
E=[];seen={}
def load(d):
 R=rows(d);h=R[0];P=col(h,'Preparer Name','Preparer');B=col(h,'BUO');L=col(h,'Lot');J=col(h,'Jockey','Jockey #1','Rider');V=col(h,'Vendor/Draft Name','Vendor');G=col(h,'Breeding');D=col(h,'Day')
 H={};nm={}
 for r in R[1:]:
  if not r[P]:continue
  p=ALIAS.get(k(r[P]),k(r[P]));nm.setdefault(p,str(r[P]).strip());j=k(r[J]);lot=r[L]
  if lot in seen:E.append(f'ERROR lot {lot} listed twice ({seen[lot]} and {d})')
  seen[lot]=d
  if j in NOJ:E.append(f'INFO {d} lot {lot} has no jockey')
  if D>=0 and r[D] and k(r[D])!=k(d):E.append(f'WARNING lot {lot} is on the {d} sheet but Day says {r[D]}')
  H.setdefault((p,float(r[B])),[]).append((r[P],r[V] if V>=0 else '',r[B],lot,r[G] if G>=0 else '',str(r[J] or 'No Jockey').strip(),'' if j in NOJ else j))
 O0=[ALIAS.get(k(r[0]),k(r[0])) for r in rows(d+' Order')[1:] if r and r[0]]
 ps={p for p,b in H};O=[p for p in O0 if p in ps]+sorted(ps-set(O0))
 for p in O:
  if p not in O0:E.append(f'WARNING {d}: {nm[p]} not on {d} Order - placed last in the preferred order')
 for x in H.values():
  js=[y[6] for y in x if y[6]]
  if len(js)>len(set(js)):E.append(f'ERROR {d}: {x[0][0]} BUO {x[0][2]} has the same jockey on both horses')
 T={}
 for p in O:
  hk=sorted(h for h in H if h[0]==p);n=PAIRS.get((p,k(d)),1);T[p]=[hk[i:i+n] for i in range(0,len(hk),n)]
 return H,O,T
def solve(H,O,T,sh,vmax,tl):
 N=len(H);m=cp.CpModel();s={h:m.NewIntVar(0,N-1,'') for h in H};m.AddAllDifferent(list(s.values()));by={};first={}
 for h,x in H.items():
  for y in x:
   if y[6]:by.setdefault(y[6],[]).append(h)
 for hs in by.values():m.AddNoOverlap([m.NewFixedSizeIntervalVar(s[h],GAP+1,'') for h in hs])
 for p in O:
  t=T[p];n=len(t);L=len(t[0]);st=[s[u[0]] for u in t]
  for u in t:
   for a,b in zip(u,u[1:]):m.Add(s[b]==s[a]+1)
  if n==1:first[p]=st[0];continue
  pos=[m.NewIntVar(0,N-1,'') for r in range(n)];S=sh//L;b={}
  for r in range(n-1):m.AddLinearConstraint(pos[r+1]-pos[r],L+VMIN,L+vmax)
  for u in range(n):
   for r in range(n):
    if abs(u-r)<=S and (len(t[u])==L or r==n-1):b[u,r]=m.NewBoolVar('');m.Add(st[u]==pos[r]).OnlyEnforceIf(b[u,r])
   m.AddExactlyOne(b[u,r] for r in range(n) if (u,r) in b)
  for r in range(n):m.AddExactlyOne(b[u,r] for u in range(n) if (u,r) in b)
  first[p]=pos[0]
 inv=[]
 for i,p in enumerate(O):
  for q in O[i+1:]:v=m.NewBoolVar('');m.Add(first[q]<first[p]).OnlyEnforceIf(v);m.Add(first[q]>=first[p]).OnlyEnforceIf(v.Not());inv.append(v)
 late=[m.NewIntVar(0,N,'') for p in O]
 for i,p in enumerate(O):m.Add(late[i]>=first[p]-i*N//len(O))
 m.Minimize(20*sum(inv)+sum(late));sv=cp.CpSolver();sv.parameters.max_time_in_seconds=tl;sv.parameters.num_workers=8
 if sv.Solve(m) not in (cp.OPTIMAL,cp.FEASIBLE):return None
 return sorted(H,key=lambda h:sv.Value(s[h])),sum(sv.Value(v) for v in inv)
out=openpyxl.load_workbook(F);summ=[];DD=[d for d in DAYS if rows(d)];T0=time.time()
for di,dn in enumerate(DD):
 H,O,T=load(dn);res=None
 for sh,vm in [(SHIFT,VMAX),(SHIFT+2,VMAX),(SHIFT,VMAX+1),(SHIFT+2,VMAX+2)]:
  res=solve(H,O,T,sh,vm,TL)
  if res:break
 if not res:summ.append(f'{dn}: no schedule found that meets all rules - check the bookings');continue
 if (sh,vm)!=(SHIFT,VMAX):E.append(f'WARNING {dn}: no schedule with BUO moves <={SHIFT} and preparer gap <={VMAX}; used BUO moves <={sh}, preparer gap <={vm}')
 seq,inv=res;ws=out.create_sheet(f'{dn} Schedule')
 ws.append(['Heat','Preparer','Vendor','BUO','Lot','Breeding','Jockey','Heats since last ride','Preparer gap','Check']);last={};pe={};jr={};jn={};ph={};cl=pg=0
 for i,h in enumerate(seq,1):
  p=h[0];q=i-pe[p]-1 if p in pe else '';pe[p]=i;ph.setdefault(p,[]).append(i);gb=q not in('',0) and not VMIN<=q<=VMAX;pg+=gb
  for x in H[h]:
   g=i-last[x[6]]-1 if x[6] in last else '';bad=g!=''and g<GAP;cl+=bad
   ws.append([i,x[0],x[1],x[2],x[3],x[4],x[5],g,q,'CLASH' if bad else 'PREP GAP' if gb else 'OK'])
   if x[6]:last[x[6]]=i;jr.setdefault(x[6],[]).append(i);jn.setdefault(x[6],x[5])
 pw=out.create_sheet(f'{dn} Preparers');pw.append(['Preferred position','Actual start position','Preparer','Heats','First heat','Last heat']);st=sorted(ph,key=lambda p:ph[p][0])
 for n,p in enumerate(O):pw.append([n+1,st.index(p)+1,H[T[p][0][0]][0][0],len(ph[p]),ph[p][0],ph[p][-1]])
 jw=out.create_sheet(f'{dn} Jockeys');jw.append(['Jockey','Rides','First heat','Last heat','Smallest gap','Longest wait',f'Waits over {JW}'])
 for j,hh in sorted(jr.items(),key=lambda z:z[1][0]-z[1][-1]):gp=[b-a-1 for a,b in zip(hh,hh[1:])];jw.append([jn[j],len(hh),hh[0],hh[-1],min(gp,default=''),max(gp,default=''),sum(g>JW for g in gp)])
 summ.append(f'{dn}: {len(seq)} heats, {cl} jockey clashes, {pg} preparer gaps outside {VMIN}-{vm}, BUO moves <={sh}, {inv} swaps vs preferred order')
v=out.create_sheet('Validation');[v.append([x]) for x in summ+E]
out.save('NZB_Breeze_Up_Schedule.xlsx');print('\n'.join(summ+[x for x in E if not x.startswith('INFO')]))
